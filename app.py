import streamlit as st
import pandas as pd
import re
import tempfile
import unicodedata
import urllib.request
import zipfile
from datetime import date
from pathlib import Path
from io import BytesIO
from fpdf import FPDF
from fpdf.enums import XPos, YPos

# =====================================================
# 1. CẤU HÌNH TRANG
# =====================================================
st.set_page_config(
    page_title="Phần mềm tính tiền lương",
    page_icon="💰",
    layout="wide"
)

# =====================================================
# 2. LOGO - KHÔNG LÀM APP BỊ LỖI NẾU THIẾU ẢNH
# =====================================================
THU_MUC_GOC = Path(__file__).resolve().parent
DUONG_DAN_LOGO = THU_MUC_GOC / "LOGO.jpg"

if DUONG_DAN_LOGO.exists():
    st.image(str(DUONG_DAN_LOGO), width=200)

# =====================================================
# 3. HÀM ĐỊNH DẠNG
# =====================================================
def dinh_dang_tien(so_tien):
    return f"{so_tien:,.3f}".replace(",", "X").replace(".", ",").replace("X", ".")


MAU_EMAIL = re.compile(r"^[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}$")


def email_hop_le(email):
    return bool(MAU_EMAIL.match(str(email).strip()))


def bo_dau(chuoi):
    # Chuyển tiếng Việt có dấu thành không dấu (dùng cho tên file / font dự phòng)
    chuoi = str(chuoi).replace("đ", "d").replace("Đ", "D")
    chuoi = unicodedata.normalize("NFD", chuoi)
    return "".join(c for c in chuoi if unicodedata.category(c) != "Mn")


def ten_file_an_toan(chuoi):
    chuoi = re.sub(r"[^A-Za-z0-9]+", "_", bo_dau(chuoi)).strip("_")
    return chuoi or "nhan_vien"

# =====================================================
# 4. HÀM TÍNH THUẾ TNCN
# Giữ nguyên biểu thuế theo code ban đầu của người dùng
# =====================================================
def tinh_thue_tncn(thu_nhap_tinh_thue):
    if thu_nhap_tinh_thue <= 0:
        return 0.0
    elif thu_nhap_tinh_thue <= 10:
        return thu_nhap_tinh_thue * 0.05
    elif thu_nhap_tinh_thue <= 30:
        return thu_nhap_tinh_thue * 0.10 - 0.5
    elif thu_nhap_tinh_thue <= 60:
        return thu_nhap_tinh_thue * 0.20 - 3.5
    elif thu_nhap_tinh_thue <= 100:
        return thu_nhap_tinh_thue * 0.30 - 9.5
    else:
        return thu_nhap_tinh_thue * 0.35 - 14.5

# =====================================================
# 5. HÀM TÍNH LƯƠNG CHO MỘT NHÂN VIÊN
# =====================================================
def tinh_luong(nhan_vien):
    luong_co_ban = float(nhan_vien["Lương cơ bản"])
    thuong = float(nhan_vien["Tiền thưởng"])
    phu_cap_co_bhxh = float(nhan_vien["Phụ cấp tính BHXH"])
    phu_cap_khong_bhxh = float(nhan_vien["Phụ cấp không tính BHXH"])
    so_nguoi_phu_thuoc = int(nhan_vien["Số người phụ thuộc"])

    tong_thu_nhap = (
        luong_co_ban
        + thuong
        + phu_cap_co_bhxh
        + phu_cap_khong_bhxh
    )

    tien_bao_hiem = (luong_co_ban + phu_cap_co_bhxh) * 0.105

    giam_tru_ban_than = 15.5
    giam_tru_nguoi_phu_thuoc = 6.2 * so_nguoi_phu_thuoc

    thu_nhap_tinh_thue = (
        tong_thu_nhap
        - giam_tru_ban_than
        - giam_tru_nguoi_phu_thuoc
        - tien_bao_hiem
    )

    thue_tncn = tinh_thue_tncn(thu_nhap_tinh_thue)
    thuc_lanh = tong_thu_nhap - tien_bao_hiem - thue_tncn

    return {
        "Tổng thu nhập": tong_thu_nhap,
        "Bảo hiểm người lao động đóng": tien_bao_hiem,
        "Giảm trừ bản thân": giam_tru_ban_than,
        "Giảm trừ người phụ thuộc": giam_tru_nguoi_phu_thuoc,
        "Thu nhập tính thuế": max(0.0, thu_nhap_tinh_thue),
        "Thuế TNCN": thue_tncn,
        "Tiền thực lãnh": thuc_lanh,
    }

# =====================================================
# 6. TẠO FILE EXCEL CÓ ĐỊNH DẠNG
# =====================================================
def tao_file_excel(df_chi_tiet):
    output = BytesIO()

    cot_tien = [
        "Lương cơ bản",
        "Tiền thưởng",
        "Phụ cấp tính BHXH",
        "Phụ cấp không tính BHXH",
        "Tổng thu nhập",
        "Bảo hiểm người lao động đóng",
        "Giảm trừ bản thân",
        "Giảm trừ người phụ thuộc",
        "Thu nhập tính thuế",
        "Thuế TNCN",
        "Tiền thực lãnh",
    ]

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        # Sheet 1: Bảng lương tổng hợp
        df_chi_tiet.to_excel(
            writer,
            sheet_name="Bang luong",
            index=False,
            startrow=2
        )

        ws = writer.sheets["Bang luong"]
        ws["A1"] = "BẢNG TỔNG HỢP LƯƠNG VÀ THUẾ TNCN"
        ws["A2"] = "Đơn vị tiền: triệu đồng/tháng"

        # Sheet 2: Chi tiết từng nhân viên
        df_chi_tiet.to_excel(
            writer,
            sheet_name="Chi tiet",
            index=False
        )

        # Định dạng cả hai sheet
        from openpyxl.styles import Font, PatternFill, Alignment

        for sheet_name in ["Bang luong", "Chi tiet"]:
            sheet = writer.sheets[sheet_name]
            header_row = 3 if sheet_name == "Bang luong" else 1

            for cell in sheet[header_row]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="1F4E78"
                )
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True
                )

            sheet.freeze_panes = f"A{header_row + 1}"
            sheet.auto_filter.ref = (
                f"A{header_row}:{sheet.cell(header_row, sheet.max_column).coordinate}"
                f"{sheet.max_row}"
            )

            for col_cells in sheet.columns:
                col_letter = col_cells[0].column_letter
                max_length = max(
                    len(str(cell.value or ""))
                    for cell in col_cells
                )
                sheet.column_dimensions[col_letter].width = min(
                    max(max_length + 2, 14), 32
                )

            for row in sheet.iter_rows(
                min_row=header_row + 1,
                max_row=sheet.max_row
            ):
                for cell in row:
                    cell.alignment = Alignment(vertical="center")
                    if cell.column > 1 and cell.value is not None:
                        # Định dạng số tiền, không định dạng cột số người phụ thuộc
                        header = sheet.cell(header_row, cell.column).value
                        if header in cot_tien:
                            cell.number_format = '#,##0.000;[Red]-#,##0.000;0.000'

        # Thêm dòng tổng cộng vào cuối bảng lương
        ws = writer.sheets["Bang luong"]
        dong_tong = ws.max_row + 1
        ws.cell(dong_tong, 1, "TỔNG CỘNG")

        cot_tong_hop = [
            "Lương cơ bản",
            "Tiền thưởng",
            "Phụ cấp tính BHXH",
            "Phụ cấp không tính BHXH",
            "Tổng thu nhập",
            "Bảo hiểm người lao động đóng",
            "Thuế TNCN",
            "Tiền thực lãnh",
        ]

        # Tìm cột theo tiêu đề, tránh phụ thuộc thứ tự cột
        tieu_de = {
            ws.cell(3, col).value: col
            for col in range(1, ws.max_column + 1)
        }

        for ten_cot in cot_tong_hop:
            if ten_cot in tieu_de:
                col = tieu_de[ten_cot]
                cot_chu = ws.cell(3, col).column_letter
                ws.cell(
                    dong_tong,
                    col,
                    f"=SUM({cot_chu}4:{cot_chu}{dong_tong - 1})"
                )
                ws.cell(dong_tong, col).number_format = '#,##0.000'

        for cell in ws[dong_tong]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="D9EAF7"
            )

    output.seek(0)
    return output.getvalue()

# =====================================================
# 7. FONT TIẾNG VIỆT CHO PDF
# Tìm font Unicode có sẵn trên máy; nếu không có (Streamlit Cloud)
# thì tự tải DejaVu Sans về thư mục tạm. Nếu vẫn không được thì
# dùng font mặc định và bỏ dấu tiếng Việt.
# =====================================================
FONT_URL = "https://cdn.jsdelivr.net/npm/dejavu-fonts-ttf@2.37.3/ttf/"


@st.cache_resource(show_spinner=False)
def tim_font_unicode():
    ung_vien = [
        (THU_MUC_GOC / "fonts" / "DejaVuSans.ttf",
         THU_MUC_GOC / "fonts" / "DejaVuSans-Bold.ttf"),
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
         Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
        (Path("C:/Windows/Fonts/arial.ttf"),
         Path("C:/Windows/Fonts/arialbd.ttf")),
    ]
    for thuong, dam in ung_vien:
        if thuong.exists() and dam.exists():
            return str(thuong), str(dam)

    thu_muc_tam = Path(tempfile.gettempdir()) / "font_phieu_luong"
    thu_muc_tam.mkdir(exist_ok=True)
    try:
        ket_qua = []
        for ten in ["DejaVuSans.ttf", "DejaVuSans-Bold.ttf"]:
            dich = thu_muc_tam / ten
            if not dich.exists():
                urllib.request.urlretrieve(FONT_URL + ten, dich)
            ket_qua.append(str(dich))
        return tuple(ket_qua)
    except Exception:
        return None

# =====================================================
# 8. TẠO PHIẾU LƯƠNG PDF CHO TỪNG NGƯỜI
# =====================================================
def tao_pdf_phieu_luong(nv, ky_luong, ten_cong_ty):
    font = tim_font_unicode()
    pdf = FPDF(format="A4")
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()

    if font:
        pdf.add_font("VN", "", font[0])
        pdf.add_font("VN", "B", font[1])
        ten_font = "VN"

        def chu(s):
            return str(s)
    else:
        ten_font = "Helvetica"

        def chu(s):
            return bo_dau(s)

    rong = pdf.w - pdf.l_margin - pdf.r_margin
    xanh = (31, 78, 120)

    # Tiêu đề
    pdf.set_font(ten_font, "B", 11)
    pdf.set_text_color(*xanh)
    pdf.cell(rong, 6, chu(ten_cong_ty.upper()),
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(4)
    pdf.set_font(ten_font, "B", 18)
    pdf.cell(rong, 10, chu("PHIẾU LƯƠNG"), align="C",
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font(ten_font, "", 11)
    pdf.set_text_color(90, 90, 90)
    pdf.cell(rong, 6, chu(f"Kỳ lương: {ky_luong}"), align="C",
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(6)

    # Thông tin nhân viên
    pdf.set_text_color(0, 0, 0)
    for nhan, gia_tri in [
        ("Họ và tên", nv["Họ và tên"]),
        ("Email", nv["Email"]),
        ("Số người phụ thuộc", int(nv["Số người phụ thuộc"])),
    ]:
        pdf.set_font(ten_font, "", 11)
        pdf.cell(50, 7, chu(f"{nhan}:"))
        pdf.set_font(ten_font, "B", 11)
        pdf.multi_cell(rong - 50, 7, chu(gia_tri),
                       new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(4)

    # Bảng chi tiết
    cot_1, cot_2 = rong * 0.68, rong * 0.32
    pdf.set_fill_color(*xanh)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font(ten_font, "B", 11)
    pdf.cell(cot_1, 9, chu("  Khoản mục"), border=1, fill=True)
    pdf.cell(cot_2, 9, chu("Số tiền (triệu đồng)  "), border=1, fill=True,
             align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    dong_bang = [
        ("Lương cơ bản", "Lương cơ bản", False),
        ("Tiền thưởng", "Tiền thưởng", False),
        ("Phụ cấp tính BHXH", "Phụ cấp tính BHXH", False),
        ("Phụ cấp không tính BHXH", "Phụ cấp không tính BHXH", False),
        ("TỔNG THU NHẬP", "Tổng thu nhập", True),
        ("Bảo hiểm người lao động đóng (10,5%)", "Bảo hiểm người lao động đóng", False),
        ("Giảm trừ bản thân", "Giảm trừ bản thân", False),
        ("Giảm trừ người phụ thuộc", "Giảm trừ người phụ thuộc", False),
        ("Thu nhập tính thuế", "Thu nhập tính thuế", False),
        ("Thuế TNCN", "Thuế TNCN", False),
    ]
    pdf.set_text_color(0, 0, 0)
    for i, (nhan, cot, dam) in enumerate(dong_bang):
        pdf.set_font(ten_font, "B" if dam else "", 11)
        to_nen = dam or i % 2 == 1
        pdf.set_fill_color(*((217, 234, 247) if dam else (245, 247, 250)))
        pdf.cell(cot_1, 8, chu(f"  {nhan}"), border=1, fill=to_nen)
        pdf.cell(cot_2, 8, dinh_dang_tien(float(nv[cot])) + "  ", border=1,
                 fill=to_nen, align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    # Thực lãnh
    pdf.set_fill_color(*xanh)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font(ten_font, "B", 13)
    pdf.cell(cot_1, 11, chu("  TIỀN THỰC LÃNH"), border=1, fill=True)
    pdf.cell(cot_2, 11, dinh_dang_tien(float(nv["Tiền thực lãnh"])) + "  ",
             border=1, fill=True, align="R",
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    # Ghi chú
    pdf.ln(6)
    pdf.set_text_color(110, 110, 110)
    pdf.set_font(ten_font, "", 9)
    pdf.multi_cell(rong, 5, chu(
        "Đơn vị tiền: triệu đồng/tháng. Tiền thực lãnh = Tổng thu nhập − "
        "Bảo hiểm − Thuế TNCN. Phiếu lương được tạo tự động ngày "
        f"{date.today():%d/%m/%Y}."
    ))

    return bytes(pdf.output())


def ten_file_pdf(stt, nv):
    return f"{stt:03d}_{ten_file_an_toan(nv['Họ và tên'])}.pdf"


def tao_zip_phieu_luong(df, ky_luong, ten_cong_ty, thanh_tien_do=None):
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as zf:
        for i, (_, nv) in enumerate(df.iterrows(), start=1):
            zf.writestr(
                ten_file_pdf(i, nv),
                tao_pdf_phieu_luong(nv, ky_luong, ten_cong_ty)
            )
            if thanh_tien_do:
                thanh_tien_do.progress(i / len(df), text=f"Đang tạo {i}/{len(df)} phiếu lương...")
        # Danh sách email kèm tên file để tiện gửi phiếu lương
        danh_sach = pd.DataFrame({
            "STT": range(1, len(df) + 1),
            "Họ và tên": df["Họ và tên"].values,
            "Email": df["Email"].values,
            "File PDF": [ten_file_pdf(i, nv) for i, (_, nv) in enumerate(df.iterrows(), start=1)],
        })
        zf.writestr("danh_sach_email.csv", danh_sach.to_csv(index=False).encode("utf-8-sig"))
    output.seek(0)
    return output.getvalue()

# =====================================================
# 9. GIAO DIỆN ỨNG DỤNG
# =====================================================
st.title("💰 PHẦN MỀM TÍNH TIỀN LƯƠNG")
st.subheader("Trần Thị Mỹ Huyền")

st.caption(
    "Tính đồng thời tiền lương, bảo hiểm, thuế TNCN và tiền thực lãnh "
    "cho tối đa 500 người lao động, xuất phiếu lương PDF cho từng người."
)

st.info(
    "Đơn vị nhập liệu là triệu đồng/tháng. "
    "Ứng dụng sử dụng các mức giảm trừ và biểu thuế từ code ban đầu. "
    "Vui lòng đối chiếu quy định pháp luật hiện hành trước khi dùng để trả lương thực tế."
)

# =====================================================
# 10. CHỌN SỐ LƯỢNG NHÂN VIÊN
# =====================================================
st.subheader("1. Chọn số lượng người lao động")

so_nhan_vien = st.number_input(
    "Số người lao động cần tính",
    min_value=1,
    max_value=500,
    value=10,
    step=1
)

st.caption(
    f"Bạn đang chọn tính lương cho {so_nhan_vien} người lao động."
)

# =====================================================
# 11. NHẬP DỮ LIỆU NHÂN VIÊN
# =====================================================
st.subheader("2. Nhập thông tin lương")

st.markdown(
    "Nhập trực tiếp trên bảng bên dưới. "
    "Có thể sao chép dữ liệu từ Excel và dán vào bảng."
)

COT_NHAP = [
    "Họ và tên",
    "Email",
    "Lương cơ bản",
    "Tiền thưởng",
    "Phụ cấp tính BHXH",
    "Phụ cấp không tính BHXH",
    "Số người phụ thuộc",
]

if "du_lieu_nhan_vien" not in st.session_state:
    st.session_state.du_lieu_nhan_vien = pd.DataFrame(
        {
            "Họ và tên": [f"Nhân viên {i + 1}" for i in range(500)],
            "Email": [f"nhanvien{i + 1}@congty.vn" for i in range(500)],
            "Lương cơ bản": [15.0] * 500,
            "Tiền thưởng": [2.0] * 500,
            "Phụ cấp tính BHXH": [1.0] * 500,
            "Phụ cấp không tính BHXH": [0.5] * 500,
            "Số người phụ thuộc": [1] * 500,
        }
    )
elif "Email" not in st.session_state.du_lieu_nhan_vien.columns:
    # Phiên cũ chưa có cột Email
    st.session_state.du_lieu_nhan_vien.insert(1, "Email", "")

# Hiển thị số dòng theo lựa chọn
du_lieu_goc = st.session_state.du_lieu_nhan_vien.iloc[:int(so_nhan_vien)][COT_NHAP].copy()

du_lieu_nhap = st.data_editor(
    du_lieu_goc,
    width="stretch",
    hide_index=True,
    num_rows="fixed",
    key=f"bang_nhap_{int(so_nhan_vien)}",
    column_config={
        "Họ và tên": st.column_config.TextColumn(
            "Họ và tên",
            required=True,
            width="medium"
        ),
        "Email": st.column_config.TextColumn(
            "Email",
            required=True,
            width="medium",
            validate=MAU_EMAIL.pattern
        ),
        "Lương cơ bản": st.column_config.NumberColumn(
            "Lương cơ bản",
            min_value=0.0,
            step=0.1,
            format="%.3f"
        ),
        "Tiền thưởng": st.column_config.NumberColumn(
            "Tiền thưởng",
            min_value=0.0,
            step=0.1,
            format="%.3f"
        ),
        "Phụ cấp tính BHXH": st.column_config.NumberColumn(
            "Phụ cấp tính BHXH",
            min_value=0.0,
            step=0.1,
            format="%.3f"
        ),
        "Phụ cấp không tính BHXH": st.column_config.NumberColumn(
            "Phụ cấp không tính BHXH",
            min_value=0.0,
            step=0.1,
            format="%.3f"
        ),
        "Số người phụ thuộc": st.column_config.NumberColumn(
            "Số người phụ thuộc",
            min_value=0,
            max_value=100,
            step=1,
            format="%d"
        ),
    }
)

# Lưu dữ liệu vừa nhập để giữ lại khi đổi số lượng
du_lieu_cap_nhat = st.session_state.du_lieu_nhan_vien.copy()
du_lieu_cap_nhat.loc[:int(so_nhan_vien) - 1, du_lieu_nhap.columns] = (
    du_lieu_nhap.reset_index(drop=True).values
)
st.session_state.du_lieu_nhan_vien = du_lieu_cap_nhat

# =====================================================
# 12. TÍNH TOÁN TOÀN BỘ NHÂN VIÊN
# =====================================================
st.divider()

if st.button(
    "🧮 TÍNH LƯƠNG TOÀN BỘ NHÂN VIÊN",
    type="primary",
    width="stretch"
):
    if du_lieu_nhap["Họ và tên"].isna().any() or (
        du_lieu_nhap["Họ và tên"].astype(str).str.strip() == ""
    ).any():
        st.error("Vui lòng nhập đầy đủ họ tên của tất cả nhân viên.")
        st.stop()

    email_sai = du_lieu_nhap[
        ~du_lieu_nhap["Email"].fillna("").map(email_hop_le)
    ]
    if not email_sai.empty:
        st.error(
            "Email chưa hợp lệ hoặc còn trống ở các nhân viên: "
            + ", ".join(email_sai["Họ và tên"].astype(str).head(10))
            + (" ..." if len(email_sai) > 10 else "")
        )
        st.stop()

    cot_so = [
        "Lương cơ bản",
        "Tiền thưởng",
        "Phụ cấp tính BHXH",
        "Phụ cấp không tính BHXH",
        "Số người phụ thuộc",
    ]

    if du_lieu_nhap[cot_so].isna().any().any():
        st.error("Vui lòng nhập đủ các số liệu lương trước khi tính.")
        st.stop()

    if (du_lieu_nhap[cot_so] < 0).any().any():
        st.error("Các khoản lương và số người phụ thuộc không được âm.")
        st.stop()

    if (
        du_lieu_nhap["Số người phụ thuộc"]
        != du_lieu_nhap["Số người phụ thuộc"].astype(int)
    ).any():
        st.error("Số người phụ thuộc phải là số nguyên.")
        st.stop()

    ket_qua = []

    for _, nv in du_lieu_nhap.iterrows():
        ket_qua_nv = tinh_luong(nv)

        dong = {
            "Họ và tên": str(nv["Họ và tên"]).strip(),
            "Email": str(nv["Email"]).strip(),
            "Lương cơ bản": float(nv["Lương cơ bản"]),
            "Tiền thưởng": float(nv["Tiền thưởng"]),
            "Phụ cấp tính BHXH": float(nv["Phụ cấp tính BHXH"]),
            "Phụ cấp không tính BHXH": float(nv["Phụ cấp không tính BHXH"]),
            "Số người phụ thuộc": int(nv["Số người phụ thuộc"]),
        }

        dong.update(ket_qua_nv)
        ket_qua.append(dong)

    df_ket_qua = pd.DataFrame(ket_qua)

    # Lưu kết quả để có thể tải file sau khi Streamlit chạy lại
    st.session_state.ket_qua_luong = df_ket_qua
    # Kết quả mới -> file ZIP cũ không còn đúng
    st.session_state.pop("file_zip_pdf", None)

# =====================================================
# 13. HIỂN THỊ KẾT QUẢ
# =====================================================
if "ket_qua_luong" in st.session_state:
    df_ket_qua = st.session_state.ket_qua_luong

    st.subheader("3. Kết quả tính lương")

    tong_thu_nhap_all = df_ket_qua["Tổng thu nhập"].sum()
    tong_bao_hiem_all = df_ket_qua["Bảo hiểm người lao động đóng"].sum()
    tong_thue_all = df_ket_qua["Thuế TNCN"].sum()
    tong_thuc_lanh_all = df_ket_qua["Tiền thực lãnh"].sum()

    c1, c2, c3, c4 = st.columns(4)

    c1.metric(
        "Số nhân viên",
        f"{len(df_ket_qua):,}"
    )

    c2.metric(
        "Tổng thu nhập",
        f"{dinh_dang_tien(tong_thu_nhap_all)} triệu"
    )

    c3.metric(
        "Tổng thuế TNCN",
        f"{dinh_dang_tien(tong_thue_all)} triệu"
    )

    c4.metric(
        "Tổng tiền thực lãnh",
        f"{dinh_dang_tien(tong_thuc_lanh_all)} triệu"
    )

    st.markdown("#### Bảng lương chi tiết")

    cot_hien_thi = [
        "Họ và tên",
        "Email",
        "Lương cơ bản",
        "Tiền thưởng",
        "Phụ cấp tính BHXH",
        "Phụ cấp không tính BHXH",
        "Số người phụ thuộc",
        "Tổng thu nhập",
        "Bảo hiểm người lao động đóng",
        "Thu nhập tính thuế",
        "Thuế TNCN",
        "Tiền thực lãnh",
    ]

    df_hien_thi = df_ket_qua[cot_hien_thi].copy()

    cot_tien_hien_thi = [
        "Lương cơ bản",
        "Tiền thưởng",
        "Phụ cấp tính BHXH",
        "Phụ cấp không tính BHXH",
        "Tổng thu nhập",
        "Bảo hiểm người lao động đóng",
        "Thu nhập tính thuế",
        "Thuế TNCN",
        "Tiền thực lãnh",
    ]

    st.dataframe(
        df_hien_thi,
        width="stretch",
        hide_index=True,
        column_config={
            cot: st.column_config.NumberColumn(
                cot,
                format="%.3f"
            )
            for cot in cot_tien_hien_thi
        }
    )

    # Dòng tổng cộng trên giao diện
    st.markdown("#### Tổng cộng toàn bộ nhân viên")

    df_tong = pd.DataFrame([{
        "Họ và tên": "TỔNG CỘNG",
        "Lương cơ bản": df_ket_qua["Lương cơ bản"].sum(),
        "Tiền thưởng": df_ket_qua["Tiền thưởng"].sum(),
        "Phụ cấp tính BHXH": df_ket_qua["Phụ cấp tính BHXH"].sum(),
        "Phụ cấp không tính BHXH": df_ket_qua["Phụ cấp không tính BHXH"].sum(),
        "Số người phụ thuộc": df_ket_qua["Số người phụ thuộc"].sum(),
        "Tổng thu nhập": tong_thu_nhap_all,
        "Bảo hiểm người lao động đóng": tong_bao_hiem_all,
        "Thu nhập tính thuế": df_ket_qua["Thu nhập tính thuế"].sum(),
        "Thuế TNCN": tong_thue_all,
        "Tiền thực lãnh": tong_thuc_lanh_all,
    }])

    st.dataframe(
        df_tong,
        width="stretch",
        hide_index=True,
        column_config={
            cot: st.column_config.NumberColumn(
                cot,
                format="%.3f"
            )
            for cot in cot_tien_hien_thi
        }
    )

    # Xuất Excel
    st.subheader("4. Xuất bảng lương Excel")

    file_excel = tao_file_excel(df_ket_qua)

    st.download_button(
        label="📥 TẢI FILE EXCEL BẢNG LƯƠNG (.xlsx)",
        data=file_excel,
        file_name="bang_luong_va_thue_TNCN.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width="stretch"
    )

    # Xuất thêm CSV nếu cần
    csv_data = df_ket_qua.to_csv(
        index=False
    ).encode("utf-8-sig")

    st.download_button(
        label="📥 TẢI FILE CSV",
        data=csv_data,
        file_name="bang_luong_va_thue_TNCN.csv",
        mime="text/csv",
        width="stretch"
    )

    # =================================================
    # 14. XUẤT PHIẾU LƯƠNG PDF CHO TỪNG NGƯỜI
    # =================================================
    st.subheader("5. Xuất phiếu lương PDF cho từng người")

    c1, c2 = st.columns(2)
    ten_cong_ty = c1.text_input("Tên công ty / đơn vị", value="CÔNG TY ABC")
    ky_luong = c2.text_input("Kỳ lương", value=f"Tháng {date.today():%m/%Y}")

    if tim_font_unicode() is None:
        st.warning(
            "Không tìm thấy font tiếng Việt và không tải được font từ Internet, "
            "PDF sẽ hiển thị tiếng Việt không dấu. Có thể đặt DejaVuSans.ttf và "
            "DejaVuSans-Bold.ttf vào thư mục 'fonts' cạnh file app."
        )

    tab_mot, tab_tat_ca = st.tabs(["📄 Từng người", "🗂️ Tất cả (ZIP)"])

    with tab_mot:
        chi_so = st.selectbox(
            "Chọn nhân viên",
            options=list(range(len(df_ket_qua))),
            format_func=lambda i: (
                f"{i + 1}. {df_ket_qua.iloc[i]['Họ và tên']} "
                f"— {df_ket_qua.iloc[i]['Email']}"
            ),
        )
        nv_chon = df_ket_qua.iloc[chi_so]

        st.download_button(
            label=f"📥 TẢI PHIẾU LƯƠNG PDF — {nv_chon['Họ và tên']}",
            data=tao_pdf_phieu_luong(nv_chon, ky_luong, ten_cong_ty),
            file_name=ten_file_pdf(chi_so + 1, nv_chon),
            mime="application/pdf",
            width="stretch"
        )

    with tab_tat_ca:
        st.caption(
            "Mỗi nhân viên một file PDF, đóng gói trong 1 file ZIP "
            "kèm danh_sach_email.csv (họ tên, email, tên file) để gửi phiếu lương."
        )

        if st.button("⚙️ TẠO PHIẾU LƯƠNG PDF CHO TẤT CẢ", width="stretch"):
            tien_do = st.progress(0.0, text="Đang tạo phiếu lương...")
            st.session_state.file_zip_pdf = tao_zip_phieu_luong(
                df_ket_qua, ky_luong, ten_cong_ty, tien_do
            )
            tien_do.empty()

        if "file_zip_pdf" in st.session_state:
            st.download_button(
                label=f"📥 TẢI FILE ZIP ({len(df_ket_qua)} PHIẾU LƯƠNG PDF)",
                data=st.session_state.file_zip_pdf,
                file_name="phieu_luong_pdf.zip",
                mime="application/zip",
                width="stretch"
            )

# =====================================================
# 15. GIẢI THÍCH VÀ LƯU Ý
# =====================================================
with st.expander("📘 Công thức tính lương"):
    st.markdown("""
    **1. Tổng thu nhập**
    Lương cơ bản + Tiền thưởng + Phụ cấp tính BHXH
    + Phụ cấp không tính BHXH.

    **2. Bảo hiểm người lao động đóng**
    (Lương cơ bản + Phụ cấp tính BHXH) × 10,5%.

    **3. Thu nhập tính thuế**
    Tổng thu nhập − 15,5 − (6,2 × số người phụ thuộc) − Bảo hiểm.

    **4. Thuế TNCN**
    Áp dụng công thức thuế lũy tiến được giữ lại từ code ban đầu.
    Nếu thu nhập tính thuế không dương, thuế được tính bằng 0.

    **5. Tiền thực lãnh**
    Tổng thu nhập − Bảo hiểm − Thuế TNCN.

    Đơn vị của các khoản tiền là triệu đồng/tháng.
    """)

st.divider()

st.caption(
    "Phần mềm tính theo các giả định được cài đặt trong mã nguồn. "
    "Cần đối chiếu quy định thuế, bảo hiểm và giảm trừ hiện hành trước khi "
    "dùng làm bảng lương chính thức."
)
